"""
Parse Excel and CSV files without pandas (openpyxl, xlrd, csv only).
Power BI users can export to Excel or CSV and upload those files.
"""
from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

import openpyxl
import xlrd


def parse_upload(content: bytes, filename: str) -> dict[str, Any]:
    """
    Parse an uploaded file (Excel or CSV) into a dict of sheet/table name -> list of dicts.
    Returns {"tables": {"Sheet1": [{"col1": v1, ...}, ...], ...}, "schema": <summary>}.
    """
    suffix = Path(filename).suffix.lower()
    tables: dict[str, list[dict[str, Any]]] = {}

    if suffix == ".xlsx":
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                tables[sheet_name] = []
            else:
                headers = [str(h) if h is not None else f"_col{i}" for i, h in enumerate(rows[0])]
                tables[sheet_name] = [
                    dict(zip(headers, (v if v is not None else "" for v in row)))
                    for row in rows[1:]
                ]
        wb.close()
    elif suffix == ".xls":
        book = xlrd.open_workbook(file_contents=content)
        for i in range(book.nsheets):
            sheet = book.sheet_by_index(i)
            rows = [sheet.row_values(j) for j in range(sheet.nrows)]
            if not rows:
                tables[sheet.name] = []
            else:
                headers = [str(h) if h else f"_col{j}" for j, h in enumerate(rows[0])]
                tables[sheet.name] = [
                    dict(zip(headers, (v if v != "" else "" for v in row)))
                    for row in rows[1:]
                ]
    elif suffix == ".csv":
        reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
        rows = list(reader)
        if not rows:
            tables["Sheet1"] = []
        else:
            headers = [str(h) if h else f"_col{i}" for i, h in enumerate(rows[0])]
            tables["Sheet1"] = [
                dict(zip(headers, (v if v else "" for v in row)))
                for row in rows[1:]
            ]
    else:
        raise ValueError(f"Unsupported format: {suffix}. Use .xlsx, .xls, or .csv")

    schema = _build_schema(tables)
    return {"tables": tables, "schema": schema}


def _build_schema(tables: dict[str, list[dict[str, Any]]]) -> str:
    parts = []
    for name, rows in tables.items():
        parts.append(f"Table: {name}")
        if not rows:
            parts.append("  Rows: 0, Columns: []")
        else:
            cols = list(rows[0].keys())
            parts.append(f"  Rows: {len(rows)}, Columns: {cols}")
        parts.append("")
    return "\n".join(parts).strip()
